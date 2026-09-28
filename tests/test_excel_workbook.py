import hashlib
import io
import json
import struct
import sys
import threading
import unittest
import warnings
import zipfile
from datetime import date, datetime, time
from decimal import Decimal, localcontext
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

import openpyxl
import openpyxl.worksheet._reader as openpyxl_reader
from openpyxl.chart import BarChart, Reference
from openpyxl.utils.cell import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src" / "analysis-agent"))

import excel_workbook as excel
from excel_workbook import ExcelError, ExcelWorkbook

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
DOC_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
MAIN_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
SHEET_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"


def _zip(parts):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, content in parts.items():
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, content)
    return buffer.getvalue()


def _xlsx(sheets):
    """Generate real OOXML entirely in memory, without openpyxl's writer temp files."""
    types = ET.Element(f"{{{CT}}}Types")
    ET.SubElement(types, f"{{{CT}}}Default", Extension="rels",
                  ContentType="application/vnd.openxmlformats-package.relationships+xml")
    ET.SubElement(types, f"{{{CT}}}Default", Extension="xml", ContentType="application/xml")
    ET.SubElement(types, f"{{{CT}}}Override", PartName="/xl/workbook.xml",
                  ContentType=MAIN_TYPE)
    root_rels = ET.Element(f"{{{REL}}}Relationships")
    ET.SubElement(root_rels, f"{{{REL}}}Relationship", Id="rId1",
                  Type=f"{DOC_REL}/officeDocument", Target="xl/workbook.xml")
    workbook = ET.Element(f"{{{NS}}}workbook")
    sheet_list = ET.SubElement(workbook, f"{{{NS}}}sheets")
    relations = ET.Element(f"{{{REL}}}Relationships")
    parts = {}
    for index, (name, rows) in enumerate(sheets.items(), 1):
        path = f"xl/worksheets/sheet{index}.xml"
        ET.SubElement(types, f"{{{CT}}}Override", PartName="/" + path,
                      ContentType=SHEET_TYPE)
        ET.SubElement(sheet_list, f"{{{NS}}}sheet", {
            "name": name, "sheetId": str(index), f"{{{DOC_REL}}}id": f"rId{index}",
        })
        ET.SubElement(relations, f"{{{REL}}}Relationship", Id=f"rId{index}",
                      Type=f"{DOC_REL}/worksheet", Target=f"worksheets/sheet{index}.xml")
        worksheet = ET.Element(f"{{{NS}}}worksheet")
        width = max((len(row) for row in rows), default=1)
        ET.SubElement(worksheet, f"{{{NS}}}dimension",
                      ref=f"A1:{get_column_letter(max(width, 1))}{max(len(rows), 1)}")
        sheet_data = ET.SubElement(worksheet, f"{{{NS}}}sheetData")
        for row_index, row_values in enumerate(rows, 1):
            row = ET.SubElement(sheet_data, f"{{{NS}}}row", r=str(row_index))
            for column_index, value in enumerate(row_values, 1):
                if value is None:
                    continue
                cell = ET.SubElement(row, f"{{{NS}}}c",
                                     r=f"{get_column_letter(column_index)}{row_index}")
                if isinstance(value, bool):
                    cell.set("t", "b")
                    ET.SubElement(cell, f"{{{NS}}}v").text = str(int(value))
                elif isinstance(value, (int, float, Decimal)):
                    cell.set("t", "n")
                    ET.SubElement(cell, f"{{{NS}}}v").text = str(value)
                elif isinstance(value, (datetime, date, time)):
                    cell.set("t", "d")
                    ET.SubElement(cell, f"{{{NS}}}v").text = value.isoformat()
                elif value.startswith("="):
                    ET.SubElement(cell, f"{{{NS}}}f").text = value[1:]
                    ET.SubElement(cell, f"{{{NS}}}v")
                else:
                    cell.set("t", "inlineStr")
                    inline = ET.SubElement(cell, f"{{{NS}}}is")
                    ET.SubElement(inline, f"{{{NS}}}t").text = value
        parts[path] = ET.tostring(worksheet)
    parts["[Content_Types].xml"] = ET.tostring(types)
    parts["_rels/.rels"] = ET.tostring(root_rels)
    parts["xl/workbook.xml"] = ET.tostring(workbook)
    parts["xl/_rels/workbook.xml.rels"] = ET.tostring(relations)
    return _zip(parts)


def _rewrite(data, replacements=None, xml_path=None, edit=None):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        parts = {info.filename: archive.read(info) for info in archive.infolist()}
    if xml_path:
        root = ET.fromstring(parts[xml_path])
        edit(root)
        parts[xml_path] = ET.tostring(root)
    parts.update(replacements or {})
    return _zip(parts)


def _cell(root, address):
    return root.find(f".//{{{NS}}}c[@r='{address}']")


def _cache(data, values, sheet=1):
    def edit(root):
        for address, (kind, value) in values.items():
            cell = _cell(root, address)
            cell.set("t", kind)
            cached = cell.find(f"{{{NS}}}v")
            if cached is None:
                cached = ET.SubElement(cell, f"{{{NS}}}v")
            cached.text = str(value)
    return _rewrite(data, xml_path=f"xl/worksheets/sheet{sheet}.xml", edit=edit)


class ExcelWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.data = _xlsx({
            "Actual": [
                ["Region", "Sales", "Label"],
                ["East", 0.1, "a"],
                ["West", 0.2, "b"],
                ["East", 0.2, "c"],
                ["East", None, "d"],
                [],
                ["West", 0.3, None],
            ],
            "Target": [
                ["Region", "Target"],
                ["East", 0.1],
                ["West", 0.4],
                ["East", 0.2],
                ["West", 0.6],
            ],
            "Empty": [],
        })
        self.book = ExcelWorkbook("sales.xlsx", self.data)

    def test_inspection_preserves_source_and_all_sheets(self):
        result = self.book.inspect()
        self.assertEqual(result["filename"], "sales.xlsx")
        self.assertEqual(result["bytes"], len(self.data))
        self.assertEqual(result["sha256"], hashlib.sha256(self.data).hexdigest())
        self.assertEqual([sheet["name"] for sheet in result["sheets"]],
                         ["Actual", "Target", "Empty"])
        actual = result["sheets"][0]
        self.assertEqual((actual["rows"], actual["columns"]), (7, 3))
        self.assertEqual(actual["headers"], ["Region", "Sales", "Label"])
        self.assertEqual(actual["formula_cells"], 0)
        self.assertEqual(actual["missing_formula_results"], 0)
        json.dumps(result, allow_nan=False)

    def test_preview_paginates_nonempty_rows_and_documents_skips(self):
        result = self.book.read_rows("Actual", offset=1, limit=2)
        self.assertEqual(result["sheet"], "Actual")
        self.assertEqual(result["columns"], ["Region", "Sales", "Label"])
        self.assertEqual(result["offset"], 1)
        self.assertEqual(result["returned_rows"], 2)
        self.assertEqual(result["total_rows"], 5)
        self.assertEqual(result["rows"], [
            {"Region": "West", "Sales": 0.2, "Label": "b"},
            {"Region": "East", "Sales": 0.2, "Label": "c"},
        ])
        self.assertTrue(result["has_more"])
        self.assertEqual(result["skipped_empty_rows"], 1)
        self.assertIn("1", " ".join(result["warnings"]))
        self.assertEqual(self.book.read_rows("Actual", offset=5)["rows"], [])
        self.assertFalse(self.book.read_rows("Actual", offset=99)["has_more"])
        last = self.book.read_rows("Actual", offset=4)
        self.assertEqual(last["rows"][0]["Label"], None)

    def test_numeric_operations_and_blank_counts(self):
        for operation, expected in [
            ("sum", 0.8), ("average", 0.2), ("min", 0.1), ("max", 0.3), ("count", 4),
        ]:
            with self.subTest(operation=operation):
                result = self.book.aggregate("Actual", "Sales", operation)
                self.assertEqual(result["rows"], 5)
                self.assertEqual(result["groups"], [
                    {"key": None, "value": expected, "count": 4, "blank_count": 1},
                ])
                self.assertEqual(result["operation"], operation)
                self.assertIsNone(result["group_by"])

    def test_grouped_aggregates_keep_first_occurrence_order(self):
        result = self.book.aggregate("Actual", "Sales", "sum", group_by="Region")
        self.assertEqual(result["groups"], [
            {"key": "East", "value": 0.3, "count": 2, "blank_count": 1},
            {"key": "West", "value": 0.5, "count": 2, "blank_count": 0},
        ])
        self.assertEqual(result["group_by"], "Region")

    def test_compare_sums_transactions_on_both_sides_and_totals(self):
        result = self.book.compare("Actual", "Target", "Region", "Sales", "Target")
        self.assertEqual(result["rows"], [
            {"key": "East", "actual": 0.3, "target": 0.3,
             "difference": 0, "attainment_pct": 100},
            {"key": "West", "actual": 0.5, "target": 1,
             "difference": -0.5, "attainment_pct": 50},
        ])
        self.assertEqual(result["totals"], {
            "actual": 0.8, "target": 1.3, "difference": -0.5, "attainment_pct": 61.5385,
        })
        self.assertEqual(result["actual_sheet"], "Actual")
        self.assertEqual(result["target_column"], "Target")
        json.dumps(result, allow_nan=False)

    def test_arithmetic_is_not_rounded_by_ambient_decimal_context(self):
        book = ExcelWorkbook("precision.xlsx", _xlsx({
            "S": [["Key", "Value"], ["a", 10**40], ["a", 0.1], ["a", -(10**40)],
                  ["a", 0.2]],
        }))
        with localcontext() as context:
            context.prec = 3
            self.assertEqual(book.aggregate("S", "Value", "sum")["groups"][0]["value"], 0.3)
            self.assertEqual(book.aggregate("S", "Value", "average")["groups"][0]["value"],
                             0.075)

    def test_dates_times_booleans_and_text_are_json_scalars(self):
        book = ExcelWorkbook("types.xlsx", _xlsx({
            "S": [["Date", "Time", "Timestamp", "Flag", "Text"],
                  [date(2026, 9, 28), time(12, 30), datetime(2026, 9, 28, 12, 30),  # noqa: DTZ001 - Excel dates have no timezone.
                   True, "memo"]],
        }))
        row = book.read_rows("S")["rows"][0]
        self.assertEqual(row, {
            "Date": "2026-09-28", "Time": "12:30:00",
            "Timestamp": "2026-09-28T12:30:00", "Flag": True, "Text": "memo",
        })
        for column in row:
            self.assertEqual(book.aggregate("S", column, "count")["groups"][0]["value"], 1)
            with self.assertRaises(ExcelError):
                book.aggregate("S", column, "sum")
        json.dumps(book.inspect(), allow_nan=False)

    def test_boolean_number_and_text_keys_do_not_collide(self):
        book = ExcelWorkbook("keys.xlsx", _xlsx({
            "S": [["Key", "Value"], [True, 10], [1, 20], ["1", 30], [1.0, 5]],
        }))
        groups = book.aggregate("S", "Value", "sum", "Key")["groups"]
        self.assertEqual(len(groups), 3)
        self.assertIs(groups[0]["key"], True)
        self.assertEqual([group["value"] for group in groups], [10, 25, 30])
        self.assertIsInstance(groups[1]["key"], (int, float))
        self.assertIsInstance(groups[2]["key"], str)

    def test_date_and_text_keys_with_same_json_value_are_ambiguous(self):
        book = ExcelWorkbook("keys.xlsx", _xlsx({
            "S": [["Key", "Value"], [date(2026, 9, 28), 1], ["2026-09-28", 2]],
        }))
        with self.assertRaisesRegex(ExcelError, "(?i)ambiguous"):
            book.aggregate("S", "Value", "sum", "Key")

    def test_cached_formulas_warn_without_touching_unrelated_missing_caches(self):
        data = _xlsx({
            "S": [["Key", "Value", "Other"], ["a", "=1+2", "=99"], ["b", "=2+3", None]],
            "Literal": [["Value"], [7]],
        })
        data = _cache(data, {"B2": ("n", 3), "B3": ("n", 5)})
        book = ExcelWorkbook("formulas.xlsx", data)
        sheet = book.inspect()["sheets"][0]
        self.assertEqual(sheet["formula_cells"], 3)
        self.assertEqual(sheet["missing_formula_results"], 1)
        self.assertRegex(" ".join(book.inspect()["warnings"]), "(?i)cached.*stale")
        self.assertEqual(book.aggregate("S", "Value", "sum")["groups"][0]["value"], 8)
        self.assertEqual(book.read_rows("S", offset=1)["rows"][0]["Value"], 5)
        self.assertEqual(book.aggregate("Literal", "Value", "sum")["groups"][0]["value"], 7)
        for action in [
            lambda: book.read_rows("S"),
            lambda: book.aggregate("S", "Other", "count"),
            lambda: book.aggregate("S", "Value", "sum", "Other"),
        ]:
            with self.assertRaisesRegex(ExcelError, "(?i)cached.*C2|C2.*cached"):
                action()

    def test_zero_formula_cache_is_not_missing_and_formula_headers_use_cache(self):
        data = _xlsx({"S": [['="Value"'], ["=0"]]})
        data = _cache(data, {"A1": ("str", "Value"), "A2": ("n", 0)})
        book = ExcelWorkbook("zero.xlsx", data)
        self.assertEqual(book.inspect()["sheets"][0]["headers"], ["Value"])
        self.assertEqual(book.aggregate("S", "Value", "sum")["groups"][0]["value"], 0)
        self.assertEqual(book.inspect()["sheets"][0]["missing_formula_results"], 0)

    def test_missing_formula_alone_is_not_skipped_as_an_empty_row(self):
        book = ExcelWorkbook("missing.xlsx", _xlsx({"S": [["Value"], ["=1"]]}))
        for action in [lambda: book.read_rows("S"),
                       lambda: book.aggregate("S", "Value", "count")]:
            with self.assertRaisesRegex(ExcelError, "(?i)cached"):
                action()

    def test_saved_empty_string_formula_results_are_blank_not_missing(self):
        data = _xlsx({"S": [["Key", "Value"], ["a", "=5"], ["b", '=IF(TRUE,"",1)'],
                            ['=""', '=IF(TRUE,"",1)']]})
        data = _cache(data, {"B2": ("n", 5)})

        def saved_empty_strings(root):
            # Excel saves a formula that returned "" as t="str" with an empty <v>.
            for address in ("B3", "A4", "B4"):
                _cell(root, address).set("t", "str")

        data = _rewrite(data, xml_path="xl/worksheets/sheet1.xml", edit=saved_empty_strings)
        book = ExcelWorkbook("blank-formulas.xlsx", data)
        sheet = book.inspect()["sheets"][0]
        self.assertEqual((sheet["formula_cells"], sheet["missing_formula_results"], sheet["data_rows"]),
                         (4, 0, 2))
        self.assertEqual(book.aggregate("S", "Value", "sum")["groups"],
                         [{"key": None, "value": 5, "count": 1, "blank_count": 1}])
        self.assertEqual(book.read_rows("S")["rows"][1], {"Key": "b", "Value": ""})

        def drop_saved_value(root):
            cell = _cell(root, "B3")
            cell.remove(cell.find(f"{{{NS}}}v"))

        missing = _rewrite(data, xml_path="xl/worksheets/sheet1.xml", edit=drop_saved_value)
        with self.assertRaisesRegex(ExcelError, "(?i)cached.*B3"):
            ExcelWorkbook("missing.xlsx", missing).aggregate("S", "Value", "sum")

    def test_totals_rows_are_included_and_flagged(self):
        data = _xlsx({
            "Sales": [["region", "revenue"], ["EMEA", 10], ["AMER", 20], ["Total", "=SUM(B2:B3)"]],
            "Targets": [["region", "target"], ["EMEA", 5], ["AMER", 5], ["Total", 10]],
            "Notes": [["item", "amount"], ["Totally new", 1], ["EMEA Total", 2], ["Grand Total", 3]],
        })
        book = ExcelWorkbook("totals.xlsx", _cache(data, {"B4": ("n", 30)}))
        result = book.aggregate("Sales", "revenue", "sum")
        self.assertEqual(result["groups"][0]["value"], 60)
        self.assertIn("row 4 ('Total')", " ".join(result["warnings"]))
        self.assertIn("double-count", " ".join(book.read_rows("Sales")["warnings"]))
        comparison = book.compare("Sales", "Targets", "region", "revenue", "target")
        self.assertEqual(comparison["totals"]["actual"], 60)
        self.assertTrue(any("'Targets'" in warning and "row 4" in warning
                            for warning in comparison["warnings"]))
        notes = " ".join(book.aggregate("Notes", "amount", "sum")["warnings"])
        self.assertIn("row 3 ('EMEA Total'), row 4 ('Grand Total')", notes)
        self.assertNotIn("Totally", notes)

    def test_error_cells_fail_only_when_accessed(self):
        data = _xlsx({"S": [["Value", "Error"], [2, "=1/0"], [3, None]]})
        data = _cache(data, {"B2": ("e", "#DIV/0!")})
        book = ExcelWorkbook("errors.xlsx", data)
        self.assertEqual(book.aggregate("S", "Value", "sum")["groups"][0]["value"], 5)
        self.assertEqual(book.inspect()["sheets"][0]["missing_formula_results"], 0)
        self.assertEqual(book.read_rows("S", offset=1)["returned_rows"], 1)
        for operation in ("count", "sum"):
            with self.assertRaisesRegex(ExcelError, r"#DIV/0!"):
                book.aggregate("S", "Error", operation)
        with self.assertRaisesRegex(ExcelError, "B2"):
            book.read_rows("S")

    def test_nonfinite_values_fail_including_count_and_preview(self):
        data = _xlsx({"S": [["Value"], [1]]})
        data = _cache(data, {"A2": ("n", "1e999")})
        book = ExcelWorkbook("infinite.xlsx", data)
        book.inspect()
        for action in [lambda: book.read_rows("S"),
                       lambda: book.aggregate("S", "Value", "count"),
                       lambda: book.aggregate("S", "Value", "sum")]:
            with self.assertRaisesRegex(ExcelError, "(?i)finite"):
                action()

    def test_blank_numeric_values_do_not_become_fake_zero(self):
        book = ExcelWorkbook("blanks.xlsx", _xlsx({
            "S": [["Key", "Value"], ["a", None], ["a", ""]],
        }))
        self.assertEqual(book.aggregate("S", "Value", "count")["groups"][0], {
            "key": None, "value": 0, "count": 0, "blank_count": 2,
        })
        for operation in ("sum", "average", "min", "max"):
            with self.assertRaisesRegex(ExcelError, "(?i)blank|numeric"):
                book.aggregate("S", "Value", operation)
        with self.assertRaises(ExcelError):
            book.compare("S", "S", "Key", "Value", "Value")

    def test_one_all_blank_group_is_null_without_discarding_other_groups(self):
        book = ExcelWorkbook("groups.xlsx", _xlsx({
            "S": [["Key", "Value"], ["a", 1], ["b", None], ["a", 2], ["b", ""]],
        }))
        for operation, value in (("sum", 3), ("average", 1.5), ("min", 1), ("max", 2)):
            with self.subTest(operation=operation):
                result = book.aggregate("S", "Value", operation, "Key")
                self.assertEqual(result["groups"], [
                    {"key": "a", "value": value, "count": 2, "blank_count": 0},
                    {"key": "b", "value": None, "count": 0, "blank_count": 2},
                ])
                self.assertIn(f"null {operation}: ['b']", " ".join(result["warnings"]))

    def test_bad_headers_and_empty_sheets_fail_operations_not_inspection(self):
        cases = [
            [["Value", "Value"], [1, 2]],
            [["Value", None], [1, 2]],
            [[None, "Value"], [1, 2]],
            [["Value", 42], [1, 2]],
            [["Value", "   "], [1, 2]],
            [[], [1, 2]],
            [],
        ]
        for rows in cases:
            with self.subTest(rows=rows):
                book = ExcelWorkbook("headers.xlsx", _xlsx({"S": rows}))
                book.inspect()
                for action in [lambda book=book: book.read_rows("S"),
                               lambda book=book: book.aggregate("S", "Value", "sum")]:
                    with self.assertRaises(ExcelError):
                        action()

    def test_empty_keys_and_key_mismatches_are_explicit_errors(self):
        for key in (None, "", "  "):
            with self.subTest(key=key):
                book = ExcelWorkbook("keys.xlsx", _xlsx({
                    "S": [["Key", "Value"], [key, 1]],
                }))
                with self.assertRaisesRegex(ExcelError, "(?i)key.*blank|blank.*key"):
                    book.aggregate("S", "Value", "sum", "Key")
        for key in ("other", True):
            with self.subTest(target_key=key):
                book = ExcelWorkbook("compare.xlsx", _xlsx({
                    "A": [["Key", "Value"], [1, 1]],
                    "T": [["Key", "Value"], [key, 1]],
                }))
                with self.assertRaisesRegex(ExcelError, "(?i)unmatched|key.*match"):
                    book.compare("A", "T", "Key", "Value", "Value")

    def test_compare_missing_key_or_value_formula_caches_fail(self):
        for actual_rows, target_rows in [
            ([["Key", "Value"], ["=1", 2]], [["Key", "Value"], [1, 2]]),
            ([["Key", "Value"], [1, 2]], [["Key", "Value"], [1, "=2"]]),
        ]:
            book = ExcelWorkbook("compare.xlsx", _xlsx({"A": actual_rows, "T": target_rows}))
            with self.assertRaisesRegex(ExcelError, "(?i)cached"):
                book.compare("A", "T", "Key", "Value", "Value")

    def test_zero_targets_have_null_ratios_and_explicit_warnings(self):
        book = ExcelWorkbook("zero.xlsx", _xlsx({
            "A": [["Key", "Value"], ["a", 2], ["b", 0]],
            "T": [["Key", "Value"], ["a", 0], ["b", 0]],
        }))
        result = book.compare("A", "T", "Key", "Value", "Value")
        self.assertTrue(all(row["attainment_pct"] is None for row in result["rows"]))
        self.assertIsNone(result["totals"]["attainment_pct"])
        self.assertEqual(result["totals"]["difference"], 2)
        warning = " ".join(result["warnings"]).lower()
        self.assertIn("zero", warning)
        self.assertIn("total", warning)

    def test_arguments_are_validated_without_coercion(self):
        for sheet in (None, 1, [], "missing"):
            with self.subTest(sheet=sheet), self.assertRaises(ExcelError):
                self.book.read_rows(sheet)
        for offset, limit in [(-1, 1), (True, 1), (1.0, 1), (0, 0), (0, 21),
                              (0, False), (0, "10")]:
            with self.subTest(offset=offset, limit=limit), self.assertRaises(ExcelError):
                self.book.read_rows("Actual", offset=offset, limit=limit)
        for column, operation, grouping in [
            ("missing", "sum", None), ([], "sum", None), ("Sales", "SUM", None),
            ("Sales", [], None), ("Sales", "sum", []), ("Sales", "sum", ""),
        ]:
            with self.subTest(column=column, operation=operation, grouping=grouping), \
                    self.assertRaises(ExcelError):
                self.book.aggregate("Actual", column, operation, grouping)

    def test_results_are_deterministic_and_do_not_expose_mutable_state(self):
        inspection = self.book.inspect()
        preview = self.book.read_rows("Actual")
        expected_inspection = self.book.inspect()
        expected_preview = self.book.read_rows("Actual")
        inspection["sheets"][0]["headers"][0] = "changed"
        preview["columns"][0] = "changed"
        preview["rows"][0]["Sales"] = 999
        self.assertEqual(self.book.inspect(), expected_inspection)
        self.assertEqual(self.book.read_rows("Actual"), expected_preview)
        self.assertEqual(ExcelWorkbook("sales.xlsx", self.data).inspect(), expected_inspection)


class ExcelWorkbookLimitsTests(unittest.TestCase):
    def setUp(self):
        self.data = _xlsx({"S": [["Key", "Value"], ["a", 1]]})

    def test_approved_limits(self):
        self.assertEqual(excel.MAX_WORKBOOK_BYTES, 5 * 1024 * 1024)
        self.assertEqual(excel.MAX_SHEETS, 20)
        self.assertEqual(excel.MAX_CELLS, 100_000)
        self.assertEqual(excel.MAX_EXPANDED_BYTES, 20 * 1024 * 1024)
        self.assertEqual(excel.MAX_PREVIEW_ROWS, 20)
        self.assertEqual(excel.MAX_GROUPS, 1000)
        self.assertLessEqual(excel.MAX_ZIP_MEMBERS, 1000)

    def test_wrong_types_extensions_corrupt_or_encrypted_input(self):
        for filename, data in [
            (None, self.data), (3, self.data), ("", self.data), ("a.xls", self.data),
            ("a.xlsm", self.data), ("a.xlsx", None), ("a.xlsx", ""),
            ("a.xlsx", bytearray(self.data)), ("a.xlsx", b""),
            ("a.xlsx", b"not a zip"), ("a.xlsx", self.data[:50]),
            ("a.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1encrypted"),
        ]:
            with self.subTest(filename=filename, data_type=type(data)), \
                    self.assertRaises(ExcelError):
                ExcelWorkbook(filename, data)

    def test_exact_file_size_limit_and_one_byte_over(self):
        with patch.object(excel, "MAX_WORKBOOK_BYTES", len(self.data)):
            ExcelWorkbook("ok.xlsx", self.data)
        with patch.object(excel, "MAX_WORKBOOK_BYTES", len(self.data) - 1), \
                self.assertRaisesRegex(ExcelError, "(?i)size|byte|limit"):
            ExcelWorkbook("over.xlsx", self.data)

    def test_sheet_count_and_workbook_total_grid_limits(self):
        data = _xlsx({"A": [["Key", "Value"], ["a", 1]],
                      "B": [["Key", "Value"], ["b", 2]]})
        with patch.object(excel, "MAX_SHEETS", 2), patch.object(excel, "MAX_CELLS", 8):
            ExcelWorkbook("ok.xlsx", data)
        for constant, value in [("MAX_SHEETS", 1), ("MAX_CELLS", 7)]:
            with self.subTest(constant=constant), patch.object(excel, constant, value), \
                    self.assertRaisesRegex(ExcelError, "(?i)limit|exceed"):
                ExcelWorkbook("over.xlsx", data)

    def test_preview_and_group_result_boundaries(self):
        data = _xlsx({"S": [["Key", "Value"]] + [[f"k{i}", i] for i in range(21)]})
        book = ExcelWorkbook("many.xlsx", data)
        self.assertEqual(book.read_rows("S", limit=20)["returned_rows"], 20)
        self.assertTrue(book.read_rows("S", limit=20)["has_more"])
        with self.assertRaises(ExcelError):
            book.read_rows("S", limit=21)
        with patch.object(excel, "MAX_GROUPS", 21):
            self.assertEqual(len(book.aggregate("S", "Value", "sum", "Key")["groups"]), 21)
        with patch.object(excel, "MAX_GROUPS", 20):
            with self.assertRaisesRegex(ExcelError, "(?i)group.*limit|exceed"):
                book.aggregate("S", "Value", "sum", "Key")
            with self.assertRaises(ExcelError):
                book.compare("S", "S", "Key", "Value", "Value")

    def test_expanded_zip_size_and_member_count_boundaries(self):
        with zipfile.ZipFile(io.BytesIO(self.data)) as archive:
            expanded_size = sum(info.file_size for info in archive.infolist())
            member_count = len(archive.infolist())
        with patch.object(excel, "MAX_EXPANDED_BYTES", expanded_size), \
                patch.object(excel, "MAX_ZIP_MEMBERS", member_count):
            ExcelWorkbook("ok.xlsx", self.data)
        for constant, value in [
            ("MAX_EXPANDED_BYTES", expanded_size - 1), ("MAX_ZIP_MEMBERS", member_count - 1),
        ]:
            with self.subTest(constant=constant), patch.object(excel, constant, value), \
                    self.assertRaisesRegex(ExcelError, "(?i)limit|exceed"):
                ExcelWorkbook("over.xlsx", self.data)

    def test_zip_bomb_style_payload_fails_before_openpyxl(self):
        data = _rewrite(self.data, {"huge.txt": b"x" * 100_000})
        self.assertLess(len(data), 5000)
        with patch.object(excel, "MAX_EXPANDED_BYTES", 20_000), \
                patch.object(excel, "load_workbook") as loader:
            with self.assertRaisesRegex(ExcelError, "(?i)expanded|uncompressed"):
                ExcelWorkbook("bomb.xlsx", data)
            loader.assert_not_called()

    def test_forged_small_dimensions_do_not_hide_rows_or_columns(self):
        data = _xlsx({"S": [["Key", "Value", "Other"], ["a", 1, "visible"],
                            [], [], ["b", 2, "last"]]})
        data = _rewrite(data, xml_path="xl/worksheets/sheet1.xml",
                        edit=lambda root: root.find(f"{{{NS}}}dimension").set("ref", "A1:A1"))
        book = ExcelWorkbook("forged.xlsx", data)
        self.assertEqual((book.inspect()["sheets"][0]["rows"],
                          book.inspect()["sheets"][0]["columns"]), (5, 3))
        self.assertEqual(book.aggregate("S", "Value", "sum")["groups"][0]["value"], 3)
        self.assertEqual(book.read_rows("S")["rows"][-1]["Other"], "last")
        with patch.object(excel, "MAX_CELLS", 14), self.assertRaises(ExcelError):
            ExcelWorkbook("over.xlsx", data)

    def test_sparse_huge_content_coordinates_fail_before_openpyxl(self):
        def huge_row(root):
            row = root.find(f".//{{{NS}}}row[@r='2']")
            row.set("r", "1048576")
            for cell in row:
                cell.set("r", cell.get("r")[:-1] + "1048576")

        edits = [
            huge_row,
            lambda root: _cell(root, "B2").set("r", "XFD2"),
            lambda root: _cell(root, "B2").set("r", "ZZZZZZZZZ2"),
        ]
        for edit in edits:
            with self.subTest(edit=edit):
                data = _rewrite(self.data, xml_path="xl/worksheets/sheet1.xml", edit=edit)
                with patch.object(excel, "MAX_CELLS", 100), \
                        patch.object(excel, "load_workbook") as loader:
                    with self.assertRaises(ExcelError):
                        ExcelWorkbook("sparse.xlsx", data)
                    loader.assert_not_called()

    def test_formatting_only_cells_hints_and_merges_do_not_size_the_grid(self):
        def formatting(root):
            sheet_data = root.find(f"{{{NS}}}sheetData")
            ET.SubElement(sheet_data.find(f"{{{NS}}}row[@r='1']"), f"{{{NS}}}c", r="C1", s="1")
            corner = ET.SubElement(sheet_data, f"{{{NS}}}row", r="1048576")
            ET.SubElement(corner, f"{{{NS}}}c", r="XFD1048576", s="1")
            root.find(f"{{{NS}}}dimension").set("ref", "A1:XFD1048576")
            merges = ET.SubElement(root, f"{{{NS}}}mergeCells", count="1")
            ET.SubElement(merges, f"{{{NS}}}mergeCell", ref="D10:Z5000")

        data = _rewrite(self.data, xml_path="xl/worksheets/sheet1.xml", edit=formatting)
        book = ExcelWorkbook("formatted.xlsx", data)
        sheet = book.inspect()["sheets"][0]
        self.assertEqual((sheet["rows"], sheet["columns"], sheet["headers"]), (2, 2, ["Key", "Value"]))
        self.assertEqual(book.aggregate("S", "Value", "sum")["groups"][0]["value"], 1)
        with patch.object(excel, "MAX_CELLS", 4):
            ExcelWorkbook("exact.xlsx", data)
        with patch.object(excel, "MAX_CELLS", 3), self.assertRaisesRegex(ExcelError, "(?i)grid limit"):
            ExcelWorkbook("over.xlsx", data)

    def test_inert_binary_parts_such_as_printer_settings_are_allowed(self):
        with zipfile.ZipFile(io.BytesIO(self.data)) as archive:
            types = ET.fromstring(archive.read("[Content_Types].xml"))
        ET.SubElement(types, f"{{{CT}}}Default", Extension="bin",
                      ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.printerSettings")
        data = _rewrite(self.data, {
            "[Content_Types].xml": ET.tostring(types),
            "xl/printerSettings/printerSettings1.bin": b"\x00" * 64,
            "xl/worksheets/_rels/sheet1.xml.rels": (
                f'<Relationships xmlns="{REL}"><Relationship Id="rId1" Type="{DOC_REL}/printerSettings" '
                'Target="../printerSettings/printerSettings1.bin"/></Relationships>'
            ).encode(),
        })
        self.assertEqual(ExcelWorkbook("print.xlsx", data).aggregate("S", "Value", "sum")["groups"][0]["value"], 1)
        with self.assertRaisesRegex(ExcelError, "(?i)macro"):
            ExcelWorkbook("macro.xlsx", _rewrite(data, {"xl/vbaProject.bin": b"macro"}))

    def test_chart_sheets_are_skipped_with_a_warning(self):
        workbook = openpyxl.Workbook()
        sales = workbook.active
        sales.title = "Sales"
        for row in [["region", "revenue"], ["EMEA", 10], ["AMER", 20]]:
            sales.append(row)
        chart = BarChart()
        chart.add_data(Reference(sales, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        workbook.create_chartsheet("Revenue chart").add_chart(chart)
        output = io.BytesIO()
        workbook.save(output)
        book = ExcelWorkbook("charts.xlsx", output.getvalue())
        result = book.inspect()
        self.assertEqual([sheet["name"] for sheet in result["sheets"]], ["Sales"])
        self.assertIn("Chart sheet 'Revenue chart' was skipped", " ".join(result["warnings"]))
        self.assertEqual(book.aggregate("Sales", "revenue", "sum")["groups"][0]["value"], 30)

    def test_corrupt_or_unsupported_zip_structures_raise_excel_error(self):
        offsets = bytearray(self.data)
        # A larger central-directory offset makes zipfile seek to negative member offsets.
        offset = struct.unpack_from("<I", offsets, len(offsets) - 6)[0]
        struct.pack_into("<I", offsets, len(offsets) - 6, offset + len(offsets))
        version = bytearray(self.data)
        # zipfile raises NotImplementedError for an unknown "version needed to extract".
        struct.pack_into("<H", version, version.find(b"PK\x01\x02") + 6, 78)
        for data in (offsets, version):
            with self.subTest(data=bytes(data[-8:])), self.assertRaisesRegex(ExcelError, "(?i)corrupt"):
                ExcelWorkbook("corrupt.xlsx", bytes(data))

    def test_concurrent_parses_keep_warnings_isolated_and_restore_global_state(self):
        original = (getattr(warnings, "_showwarnmsg_impl", None), warnings.filters)
        a_entered, b_entered, a_done = threading.Event(), threading.Event(), threading.Event()
        real_load, books = excel.load_workbook, {}

        def load(buffer, **kwargs):
            if not kwargs["data_only"]:
                if threading.current_thread().name == "A":
                    a_entered.set()
                    b_entered.wait(1)
                    warnings.warn_explicit("warning from workbook A", UserWarning, openpyxl_reader.__file__, 1)
                    warnings.warn("unrelated warning from another component")
                else:
                    b_entered.set()
                    a_done.wait(1)
            return real_load(buffer, **kwargs)

        def parse(name):
            books[name] = ExcelWorkbook(f"{name}.xlsx", self.data)
            if name == "A":
                a_done.set()

        with patch.object(excel, "load_workbook", side_effect=load):
            first = threading.Thread(target=parse, args=("A",), name="A")
            second = threading.Thread(target=parse, args=("B",), name="B")
            first.start()
            a_entered.wait(5)
            second.start()
            first.join(10)
            second.join(10)
        self.assertIn("warning from workbook A", books["A"].inspect()["warnings"])
        self.assertNotIn("warning from workbook A", books["B"].inspect()["warnings"])
        self.assertNotIn("unrelated warning from another component", books["A"].inspect()["warnings"])
        self.assertIs(getattr(warnings, "_showwarnmsg_impl", None), original[0])
        self.assertIs(warnings.filters, original[1])

    def test_duplicate_or_out_of_order_rows_and_cells_cannot_be_silently_lost(self):
        def duplicate_row(root):
            root.find(f".//{{{NS}}}row[@r='2']").set("r", "1")

        def duplicate_cell(root):
            _cell(root, "B2").set("r", "A2")

        def mismatch(root):
            _cell(root, "B2").set("r", "B3")

        def reverse_rows(root):
            rows = root.find(f"{{{NS}}}sheetData")
            rows[:] = reversed(list(rows))

        for edit in (duplicate_row, duplicate_cell, mismatch, reverse_rows):
            with self.subTest(edit=edit), self.assertRaises(ExcelError):
                ExcelWorkbook("bad.xlsx", _rewrite(
                    self.data, xml_path="xl/worksheets/sheet1.xml", edit=edit))

    def test_malformed_and_unsafe_xml_macros_and_unsafe_member_names(self):
        for replacements in [
            {"xl/worksheets/sheet1.xml": b"<worksheet"},
            {"extra.xml": b'<!DOCTYPE x [<!ENTITY e "bad">]><x>&e;</x>'},
            {"xl/vbaProject.bin": b"macro"},
            {"../outside.xml": b"<x/>"},
        ]:
            with self.subTest(parts=list(replacements)), self.assertRaises(ExcelError):
                ExcelWorkbook("bad.xlsx", _rewrite(self.data, replacements))
        data = self.data
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            content_types = archive.read("[Content_Types].xml").replace(
                MAIN_TYPE.encode(), b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")
        with self.assertRaisesRegex(ExcelError, "(?i)macro|unsupported"):
            ExcelWorkbook("disguised.xlsx", _rewrite(
                self.data, {"[Content_Types].xml": content_types}))

    def test_encrypted_zip_members_are_rejected_explicitly(self):
        data = bytearray(self.data)
        for signature, flag_offset in [(b"PK\x03\x04", 6), (b"PK\x01\x02", 8)]:
            position = data.find(signature)
            flags = struct.unpack_from("<H", data, position + flag_offset)[0]
            struct.pack_into("<H", data, position + flag_offset, flags | 1)
        with self.assertRaisesRegex(ExcelError, "(?i)encrypt"):
            ExcelWorkbook("encrypted.xlsx", bytes(data))

    def test_both_workbooks_and_buffers_are_closed_on_success_and_parse_failure(self):
        bad = _cache(self.data, {"B2": ("n", "not-a-number")})
        for data, should_fail in [(self.data, False), (bad, True)]:
            workbooks, buffers, options = [], [], []

            def load(buffer, buffers=buffers, options=options, workbooks=workbooks, **kwargs):
                buffers.append(buffer)
                options.append(kwargs)
                workbook = openpyxl.load_workbook(buffer, **kwargs)
                workbooks.append(workbook)
                return workbook

            with self.subTest(should_fail=should_fail), \
                    patch.object(excel, "load_workbook", side_effect=load):
                if should_fail:
                    with self.assertRaises(ExcelError):
                        ExcelWorkbook("bad.xlsx", data)
                else:
                    ExcelWorkbook("ok.xlsx", data)
            self.assertEqual(len(workbooks), 2)
            self.assertTrue(all(book._archive.fp is None for book in workbooks))
            self.assertTrue(all(buffer.closed for buffer in buffers))
            self.assertEqual({option["data_only"] for option in options}, {True, False})
            self.assertTrue(all(option["read_only"] and not option["keep_links"]
                                and not option["keep_vba"] for option in options))

    def test_first_workbook_is_closed_if_second_load_fails(self):
        workbooks, buffers = [], []

        def load(buffer, **kwargs):
            buffers.append(buffer)
            if workbooks:
                raise ValueError("invalid second workbook")
            workbook = openpyxl.load_workbook(buffer, **kwargs)
            workbooks.append(workbook)
            return workbook

        with patch.object(excel, "load_workbook", side_effect=load), \
                self.assertRaises(ExcelError):
            ExcelWorkbook("bad.xlsx", self.data)
        self.assertIsNone(workbooks[0]._archive.fp)
        self.assertTrue(all(buffer.closed for buffer in buffers))

    def test_no_filesystem_reads_or_extraction(self):
        with patch("builtins.open", side_effect=AssertionError("filesystem access")), \
                patch.object(zipfile.ZipFile, "extract",
                             side_effect=AssertionError("extraction")), \
                patch.object(zipfile.ZipFile, "extractall",
                             side_effect=AssertionError("extraction")):
            book = ExcelWorkbook("local.xlsx", self.data)
            book.inspect()
            book.read_rows("S")
            book.aggregate("S", "Value", "sum")
            book.compare("S", "S", "Key", "Value", "Value")


if __name__ == "__main__":
    unittest.main()
